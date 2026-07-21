#!/usr/bin/env python3
"""Gate: spreadsheets (MILESTONE v1.9 Wave 4, pack B). Gates .xlsx deliverables.

Canonical v2 output (ADR-SEALED-GATES decision 3): one `FAIL <ID> <category>` line per
failing check, then `gate: N/5`. Exit 0 iff all pass. Fail-closed on any crash.

Usage: gate.py --config <config.json> <artifact.xlsx>   (artifact appended by gate-runner)

Config shape (the card declares these ranges):
  {
    "formula_ranges": ["Data!B2:B11", "Summary!B2"],
    "perturbation": {"input_cell": "Inputs!B2", "observe_cell": "Summary!B2", "delta": 13},
    "totals": [{"range": "Data!B2:B10", "total_cell": "Data!B11", "tolerance": 0.01}]
  }

Checks (complete inventory, mirrored in card.md):
  SS-01 execution  workbook opens and recalculates headless with zero error cells
                   (#REF!, #DIV/0!, #VALUE!, #NAME?). Primary engine: the `formulas`
                   pip library (real dependency-graph recalc). DOCUMENTED DEGRADATION:
                   when `formulas` is not importable (or FERROX_SS_NO_RECALC=1), the
                   check degrades to an error scan of cached values plus error tokens
                   embedded in formula strings; openpyxl alone cannot recalculate.
  SS-02 value      every cell in the declared formula_ranges contains a formula, not a
                   pasted literal.
  SS-03 relation   perturbation probe: mutating the declared input cell MUST change the
                   declared observed total (blocks hardcoded totals). Dynamic when
                   recalc is available (input override, observed value must move);
                   degraded mode falls back to a static dependency walk: the observed
                   cell must be a formula whose transitive reference closure reaches
                   the input cell.
  SS-04 relation   every cross-sheet reference in every formula names a sheet that
                   exists in the workbook.
  SS-05 relation   declared totals recompute: the sum over each declared range must
                   match its total cell within tolerance (computed values when recalc
                   is available; cached values when present; otherwise the total cell
                   formula must be a SUM over exactly the declared range).

Runtime: python3 + openpyxl (required), formulas (optional recalc engine; EUPL-1.1
license is outside the shipped-pack allowlist, so it stays an optional host dependency
declared on the card, never vendored). Orchestrator-authored; fail-closed.
"""

import contextlib
import io
import json
import os
import re
import sys

CHECKS = [
    ("SS-01", "execution"),
    ("SS-02", "value"),
    ("SS-03", "relation"),
    ("SS-04", "relation"),
    ("SS-05", "relation"),
]

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
RANGE_EXPANSION_CAP = 20000

QUOTED_SHEET_RE = re.compile(r"'([^']+)'!")
BARE_SHEET_RE = re.compile(r"(?<![A-Za-z0-9_.'\]#])([A-Za-z_][A-Za-z0-9_.]*)!")
CELL_REF_RE = re.compile(
    r"(?:('(?:[^']+)'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(\$?[A-Z]{1,3}\$?[0-9]+(?::\$?[A-Z]{1,3}\$?[0-9]+)?)(?!\()"
)
SUM_ONLY_RE = re.compile(
    r"^=\s*SUM\(\s*(?:('(?:[^']+)'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(\$?[A-Z]{1,3}\$?[0-9]+:\$?[A-Z]{1,3}\$?[0-9]+)\s*\)\s*$",
    re.IGNORECASE,
)


def fail_closed(msg=None):
    for check_id, category in CHECKS:
        print("FAIL %s %s" % (check_id, category))
    print("gate: 0/%d" % len(CHECKS))
    sys.exit(1)


def col_to_num(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n


def num_to_col(n):
    out = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def parse_ref(ref, default_sheet=None):
    """'Sheet!A1' or 'Sheet!A1:B10' or bare 'A1' -> (SHEET, ((c1,r1),(c2,r2)))."""
    ref = ref.replace("$", "").strip()
    sheet = default_sheet
    if "!" in ref:
        sheet_part, ref = ref.rsplit("!", 1)
        sheet = sheet_part.strip("'")
    m = re.match(r"^([A-Z]{1,3})([0-9]+)(?::([A-Z]{1,3})([0-9]+))?$", ref.upper())
    if m is None or sheet is None:
        return None
    c1, r1 = col_to_num(m.group(1)), int(m.group(2))
    if m.group(3) is not None:
        c2, r2 = col_to_num(m.group(3)), int(m.group(4))
    else:
        c2, r2 = c1, r1
    return (sheet.upper(), ((min(c1, c2), min(r1, r2)), (max(c1, c2), max(r1, r2))))


def expand(parsed):
    """(SHEET, box) -> [(SHEET, 'A1'), ...] with an expansion cap."""
    sheet, ((c1, r1), (c2, r2)) = parsed
    cells = []
    count = 0
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            count += 1
            if count > RANGE_EXPANSION_CAP:
                return cells
            cells.append((sheet, "%s%d" % (num_to_col(c), r)))
    return cells


def is_formula(cell):
    if getattr(cell, "data_type", None) == "f":
        return True
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        return True
    return v is not None and v.__class__.__name__ == "ArrayFormula"


def formula_text(cell):
    v = cell.value
    if isinstance(v, str):
        return v
    text = getattr(v, "text", None)
    return text if isinstance(text, str) else ""


def collect_formulas(wb):
    """{(SHEET, CELL): formula_string} across every sheet."""
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if is_formula(cell):
                    out[(ws.title.upper(), cell.coordinate)] = formula_text(cell)
    return out


def formula_refs(formula, default_sheet):
    """All (SHEET, CELL) references a formula string touches, ranges expanded."""
    refs = []
    for sheet_part, ref in CELL_REF_RE.findall(formula):
        sheet = sheet_part.strip("'") if sheet_part else default_sheet
        parsed = parse_ref("%s!%s" % (sheet, ref)) if sheet else None
        if parsed is not None:
            refs.extend(expand(parsed))
    return refs


def try_recalc(path):
    """Real headless recalc via the formulas library. Returns {computed, model} where
    computed maps (SHEET, CELL) -> scalar. None when the engine is unavailable."""
    if os.environ.get("FERROX_SS_NO_RECALC") == "1":
        return None
    try:
        import formulas  # noqa: F401
    except Exception:
        return None
    try:
        sink = io.StringIO()
        with contextlib.redirect_stdout(sink):
            model = formulas.ExcelModel().loads(path).finish()
            solution = model.calculate()
        computed = {}
        key_re = re.compile(r"^'?\[[^\]]*\]([^']+)'?!(.+)$")
        for key, wrapped in solution.items():
            m = key_re.match(key)
            if m is None:
                continue
            value = getattr(wrapped, "value", wrapped)
            try:
                flat = value.ravel().tolist()
            except Exception:
                flat = [value]
            if len(flat) != 1:
                continue
            computed[(m.group(1).upper(), m.group(2).upper())] = flat[0]
        return {"computed": computed, "model": model}
    except Exception:
        return {"computed": None, "model": None, "raised": True}


def scalar_number(value):
    try:
        if isinstance(value, bool) or value is None:
            return None
        return float(value)
    except Exception:
        return None


def main():
    args = sys.argv[1:]
    config_path = None
    positional = []
    i = 0
    while i < len(args):
        if args[i] == "--config":
            i += 1
            config_path = args[i] if i < len(args) else None
        else:
            positional.append(args[i])
        i += 1
    if not positional:
        fail_closed()
    artifact = positional[-1]

    # Sealed-store runs hand the gate an extensionless artifact path; openpyxl and the
    # formulas loader are both extension-driven, so shadow-copy to a .xlsx name first.
    if not artifact.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        import atexit
        import shutil
        import tempfile

        shadow_dir = tempfile.mkdtemp(prefix="ss-gate-artifact-")
        atexit.register(shutil.rmtree, shadow_dir, ignore_errors=True)
        shadow = os.path.join(shadow_dir, "artifact.xlsx")
        try:
            shutil.copyfile(artifact, shadow)
        except Exception:
            fail_closed()
        artifact = shadow

    config = {}
    if config_path:
        try:
            with open(config_path, "r", encoding="utf-8") as fh:
                config = json.load(fh)
        except Exception:
            config = {}

    try:
        import openpyxl
    except Exception:
        fail_closed()

    try:
        wb = openpyxl.load_workbook(artifact, data_only=False)
        wb_cached = openpyxl.load_workbook(artifact, data_only=True)
    except Exception:
        fail_closed()

    sheet_names = {name.upper() for name in wb.sheetnames}
    formulas_map = collect_formulas(wb)

    recalc = try_recalc(artifact)
    recalc_raised = bool(recalc and recalc.get("raised"))
    computed = recalc["computed"] if recalc and recalc.get("computed") else None
    model = recalc["model"] if recalc else None

    results = {}

    # SS-01: opens + recalculates with zero error cells.
    ss01_ok = True
    if recalc_raised:
        ss01_ok = False
    if computed is not None:
        for value in computed.values():
            if isinstance(value, str) and any(tok in value for tok in ERROR_TOKENS):
                ss01_ok = False
                break
    else:
        for ws in wb_cached.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(tok in cell.value for tok in ERROR_TOKENS):
                        ss01_ok = False
    if ss01_ok:
        for formula in formulas_map.values():
            if any(tok in formula for tok in ERROR_TOKENS):
                ss01_ok = False
                break
    results["SS-01"] = ss01_ok

    # SS-02: declared formula ranges hold formulas, not literals.
    ss02_ok = True
    for range_ref in config.get("formula_ranges", []):
        parsed = parse_ref(range_ref)
        if parsed is None:
            ss02_ok = False
            continue
        for target in expand(parsed):
            if target not in formulas_map:
                ss02_ok = False
    results["SS-02"] = ss02_ok

    # SS-03: perturbation probe.
    ss03_ok = True
    perturb = config.get("perturbation") or {}
    input_ref = parse_ref(perturb.get("input_cell", "")) if perturb.get("input_cell") else None
    observe_ref = parse_ref(perturb.get("observe_cell", "")) if perturb.get("observe_cell") else None
    if input_ref is not None and observe_ref is not None:
        input_key = expand(input_ref)[0]
        observe_key = expand(observe_ref)[0]
        delta = perturb.get("delta", 1)
        dynamic_done = False
        if computed is not None and model is not None:
            base_input = scalar_number(computed.get(input_key))
            base_observe = scalar_number(computed.get(observe_key))
            if base_input is not None:
                try:
                    solution_key = None
                    sink = io.StringIO()
                    with contextlib.redirect_stdout(sink):
                        base = model.calculate()
                        key_re = re.compile(r"^'?\[[^\]]*\]([^']+)'?!(.+)$")
                        for key in base:
                            m = key_re.match(key)
                            if m and (m.group(1).upper(), m.group(2).upper()) == input_key:
                                solution_key = key
                                break
                        if solution_key is not None:
                            perturbed = model.calculate(inputs={solution_key: base_input + delta})
                    if solution_key is not None:
                        new_observe = None
                        for key, wrapped in perturbed.items():
                            m = key_re.match(key)
                            if m and (m.group(1).upper(), m.group(2).upper()) == observe_key:
                                value = getattr(wrapped, "value", wrapped)
                                try:
                                    new_observe = value.ravel().tolist()[0]
                                except Exception:
                                    new_observe = value
                                break
                        new_num = scalar_number(new_observe)
                        if base_observe is None or new_num is None or abs(new_num - base_observe) <= 1e-9:
                            ss03_ok = False
                        dynamic_done = True
                except Exception:
                    ss03_ok = False
                    dynamic_done = True
        if not dynamic_done:
            # Static dependency walk: observe must be a formula whose closure reaches input.
            start = observe_key
            if start not in formulas_map:
                ss03_ok = False
            else:
                seen = set()
                stack = [start]
                reached = False
                while stack and not reached:
                    node = stack.pop()
                    if node in seen:
                        continue
                    seen.add(node)
                    formula = formulas_map.get(node)
                    if formula is None:
                        continue
                    for ref in formula_refs(formula, node[0]):
                        if ref == input_key:
                            reached = True
                            break
                        if ref in formulas_map and ref not in seen:
                            stack.append(ref)
                if not reached:
                    ss03_ok = False
    results["SS-03"] = ss03_ok

    # SS-04: every cross-sheet reference resolves to an existing sheet.
    ss04_ok = True
    for formula in formulas_map.values():
        for name in QUOTED_SHEET_RE.findall(formula):
            if name.upper() not in sheet_names:
                ss04_ok = False
        for name in BARE_SHEET_RE.findall(formula):
            if name.upper() not in sheet_names:
                ss04_ok = False
    results["SS-04"] = ss04_ok

    # SS-05: declared totals recompute within tolerance.
    ss05_ok = True
    for spec in config.get("totals", []):
        parsed_range = parse_ref(spec.get("range", ""))
        parsed_total = parse_ref(spec.get("total_cell", ""))
        if parsed_range is None or parsed_total is None:
            ss05_ok = False
            continue
        tolerance = float(spec.get("tolerance", 0.01))
        range_cells = expand(parsed_range)
        total_key = expand(parsed_total)[0]

        def cell_value(key, source):
            if source == "computed":
                return scalar_number(computed.get(key))
            ws = None
            for candidate in wb_cached.worksheets:
                if candidate.title.upper() == key[0]:
                    ws = candidate
                    break
            return scalar_number(ws[key[1]].value) if ws is not None else None

        verified = False
        for source in (("computed",) if computed is not None else ()) + ("cached",):
            total_value = cell_value(total_key, source)
            if total_value is None:
                continue
            summed = 0.0
            usable = True
            for key in range_cells:
                value = cell_value(key, source)
                if value is None:
                    usable = False
                    break
                summed += value
            if not usable:
                continue
            if abs(summed - total_value) > tolerance:
                ss05_ok = False
            verified = True
            break
        if not verified:
            # Static fallback: the total formula must be SUM over exactly the declared range.
            formula = formulas_map.get(total_key)
            if formula is None:
                continue  # literal total with unverifiable range: SS-02/SS-03 own that mode
            m = SUM_ONLY_RE.match(formula.strip())
            if m is None:
                ss05_ok = False
                continue
            sheet_part = (m.group(1) or total_key[0]).strip("'")
            declared = parse_ref("%s!%s" % (sheet_part, m.group(2)))
            if declared != parsed_range:
                ss05_ok = False
    results["SS-05"] = ss05_ok

    passed = 0
    for check_id, category in CHECKS:
        if results.get(check_id, False):
            passed += 1
        else:
            print("FAIL %s %s" % (check_id, category))
    print("gate: %d/%d" % (passed, len(CHECKS)))
    sys.exit(0 if passed == len(CHECKS) else 1)


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception:
        fail_closed()
