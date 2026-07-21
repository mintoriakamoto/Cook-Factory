#!/usr/bin/env python3
"""Fluent fixture generators for the spreadsheets gate (v1.9 Wave 4, pack B).

ADR-SEALED-GATES: fixture CONTENT never enters the repo. Only this generator is
committed; the orchestrator (and the test suite) runs it, seals the emitted workbooks
into the sealed store, and fills the card's `sealed:sha256:` references at seal time.
Every workbook carries a per-seal nonce cell so sealed instances differ per machine
and can never hash-collide with a repo blob.

Usage: generators.py --out <dir> [--nonce <hex>]

Emits into <dir>:
  reference.xlsx           passes all 5 checks
  mutant-ss-m1.xlsx .. m5  the fluent pool
  config.json              the gate config the card invocation passes via --config
  pool.json                [{id, file, why_fluent, expected_drop, must_fail}]

The model: Inputs!B2 unit price feeds Data!B2:B10 line totals (=A{r}*Inputs!B2),
Data!B11 sums them, Summary!B2 mirrors the grand total cross-sheet, and Archive is the
rarely-viewed sheet every real workbook drags along.

Pool (5 mutants, all fluent-but-wrong):
  ss-m1  hardcoded grand total equal to the true sum        must_fail SS-02, SS-03
  ss-m2  cross-sheet ref to a plausible missing sheet       must_fail SS-04
  ss-m3  SUM range off by 1 row (drops the last data row)   must_fail SS-05
  ss-m4  #REF! buried in the rarely-viewed Archive sheet    must_fail SS-01
  ss-m5  literals pasted over formulas in half the range    must_fail SS-02
"""

import argparse
import json
import os
import secrets

import openpyxl

QUANTITIES = {r: r for r in range(2, 11)}  # row -> quantity
UNIT_PRICE = 40


def build_workbook(nonce):
    wb = openpyxl.Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"] = "Assumption"
    inputs["B1"] = "Value"
    inputs["A2"] = "Unit price"
    inputs["B2"] = UNIT_PRICE
    inputs["D1"] = "batch-%s" % nonce

    data = wb.create_sheet("Data")
    data["A1"] = "Quantity"
    data["B1"] = "Line total"
    for r, qty in QUANTITIES.items():
        data.cell(row=r, column=1, value=qty)
        data.cell(row=r, column=2, value="=A%d*Inputs!B2" % r)
    data["A11"] = "Grand total"
    data["B11"] = "=SUM(B2:B10)"

    summary = wb.create_sheet("Summary")
    summary["A2"] = "Grand total"
    summary["B2"] = "=Data!B11"

    archive = wb.create_sheet("Archive")
    archive["A1"] = "Archived 2025 run"
    archive["C7"] = 1988

    return wb


def true_total():
    return sum(QUANTITIES.values()) * UNIT_PRICE


def mutate_m1(wb):
    """Hardcoded total that looks right: the literal equals the true sum today."""
    wb["Data"]["B11"] = true_total()


def mutate_m2(wb):
    """Cross-sheet ref to Assumptions, a plausible sheet name that does not exist."""
    wb["Summary"]["B2"] = "='Assumptions'!B11"


def mutate_m3(wb):
    """SUM range excludes the last data row."""
    wb["Data"]["B11"] = "=SUM(B2:B9)"


def mutate_m4(wb):
    """#REF! buried in the rarely-viewed Archive sheet."""
    wb["Archive"]["C7"] = "=#REF!+1"


def mutate_m5(wb):
    """Literals pasted over formulas in half the declared range; values still correct."""
    data = wb["Data"]
    for r in range(2, 6):
        data.cell(row=r, column=2, value=QUANTITIES[r] * UNIT_PRICE)


POOL = [
    {
        "id": "ss-m1",
        "mutate": mutate_m1,
        "why_fluent": "the grand total literal equals the true sum today, so every rendered number is right; only formula-ness and perturbation expose it",
        "expected_drop": 2,
        "must_fail": ["SS-02", "SS-03"],
    },
    {
        "id": "ss-m2",
        "mutate": mutate_m2,
        "why_fluent": "references an Assumptions sheet, the most plausible name a finance workbook could carry; the formula reads clean",
        "expected_drop": 1,
        "must_fail": ["SS-04"],
    },
    {
        "id": "ss-m3",
        "mutate": mutate_m3,
        "why_fluent": "SUM(B2:B9) drops only the last data row; the visible shape of the sheet is untouched and the number still looks plausible",
        "expected_drop": 1,
        "must_fail": ["SS-05"],
    },
    {
        "id": "ss-m4",
        "mutate": mutate_m4,
        "why_fluent": "the #REF! sits in the Archive sheet nobody opens; all headline sheets render perfectly",
        "expected_drop": 1,
        "must_fail": ["SS-01"],
    },
    {
        "id": "ss-m5",
        "mutate": mutate_m5,
        "why_fluent": "half the line-total column is pasted literals whose values are currently correct, so the sheet renders identical to the real one",
        "expected_drop": 1,
        "must_fail": ["SS-02"],
    },
]


def gate_config():
    return {
        "formula_ranges": ["Data!B2:B11", "Summary!B2"],
        "perturbation": {"input_cell": "Inputs!B2", "observe_cell": "Summary!B2", "delta": 13},
        "totals": [{"range": "Data!B2:B10", "total_cell": "Data!B11", "tolerance": 0.01}],
    }


def emit(out_dir, nonce):
    os.makedirs(out_dir, exist_ok=True)
    reference = build_workbook(nonce)
    reference.save(os.path.join(out_dir, "reference.xlsx"))

    pool_manifest = []
    for entry in POOL:
        wb = build_workbook(nonce)
        entry["mutate"](wb)
        filename = "mutant-%s.xlsx" % entry["id"]
        wb.save(os.path.join(out_dir, filename))
        pool_manifest.append(
            {
                "id": entry["id"],
                "file": filename,
                "why_fluent": entry["why_fluent"],
                "expected_drop": entry["expected_drop"],
                "must_fail": entry["must_fail"],
            }
        )

    with open(os.path.join(out_dir, "config.json"), "w", encoding="utf-8") as fh:
        json.dump(gate_config(), fh, indent=2)
    with open(os.path.join(out_dir, "pool.json"), "w", encoding="utf-8") as fh:
        json.dump(pool_manifest, fh, indent=2)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    parser.add_argument("--nonce", default=None)
    args = parser.parse_args()
    nonce = args.nonce if args.nonce else secrets.token_hex(4)
    emit(args.out, nonce)
    print("emitted reference + %d mutants to %s" % (len(POOL), args.out))


if __name__ == "__main__":
    main()
